"""Download pages 1-7 of https://dge.gov.in/nat and save them as an Excel sheet.

Setup (once):
    pip install requests lxml openpyxl

Run:
    python scrape_dge_nat.py
    python scrape_dge_nat.py --pages 7 --out dge_nat.xlsx

The site numbers pages from 0, so "page 1" is ?page=0 and "page 7" is ?page=6.
"""

import argparse
import sys
import time
from urllib.parse import urljoin

import requests
from lxml import html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://dge.gov.in/nat"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"
    ),
    "Accept-Language": "en-IN,en;q=0.9",
}


def clean(text):
    return " ".join((text or "").split())


def fetch(session, url, retries=4):
    for attempt in range(1, retries + 1):
        try:
            resp = session.get(url, headers=HEADERS, timeout=60)
            resp.raise_for_status()
            return resp.text
        except requests.RequestException as exc:
            if attempt == retries:
                raise
            wait = 2 ** attempt
            print(f"  retry {attempt}/{retries - 1} in {wait}s ({exc})")
            time.sleep(wait)


def cell_links(el, page_url):
    hrefs = []
    for a in el.xpath(".//a[@href]"):
        href = a.get("href").strip()
        if href and not href.startswith(("#", "javascript:")):
            hrefs.append(urljoin(page_url, href))
    return hrefs


def parse_table(doc, page_url):
    """Return (headers, rows) from the biggest <table> on the page."""
    tables = doc.xpath("//table")
    if not tables:
        return None
    table = max(tables, key=lambda t: len(t.xpath(".//tr")))

    headers = [clean(th.text_content()) for th in table.xpath(".//thead//th")]
    body_rows = table.xpath(".//tbody/tr") or table.xpath(".//tr")
    if not headers and body_rows and body_rows[0].xpath("./th"):
        headers = [clean(th.text_content()) for th in body_rows[0].xpath("./th")]
        body_rows = body_rows[1:]

    rows = []
    for tr in body_rows:
        cells = tr.xpath("./td|./th")
        if not cells:
            continue
        values = [clean(c.text_content()) for c in cells]
        links = [link for c in cells for link in cell_links(c, page_url)]
        rows.append(values + [" | ".join(dict.fromkeys(links))])
    return headers, rows


def parse_views_rows(doc, page_url):
    """Fallback for Drupal listings laid out as <div class="views-row">."""
    items = doc.xpath('//div[contains(concat(" ", @class, " "), " views-row ")]')
    if not items:
        return None
    headers, rows = [], []
    for item in items:
        fields = item.xpath('.//*[contains(@class, "views-field")]') or [item]
        if not headers:
            for f in fields:
                label = f.xpath('.//*[contains(@class, "views-label")]')
                headers.append(clean(label[0].text_content()).rstrip(":") if label else "")
        values = []
        for f in fields:
            content = f.xpath('.//*[contains(@class, "field-content")]')
            values.append(clean((content[0] if content else f).text_content()))
        rows.append(values + [" | ".join(dict.fromkeys(cell_links(item, page_url)))])
    return headers, rows


def scrape(pages):
    session = requests.Session()
    all_headers, all_rows = [], []
    for n in range(pages):
        url = f"{BASE_URL}?page={n}"
        print(f"Page {n + 1}: {url}")
        doc = html.fromstring(fetch(session, url))
        result = parse_table(doc, url) or parse_views_rows(doc, url)
        if not result or not result[1]:
            print("  no data found on this page, stopping")
            break
        headers, rows = result
        if len(headers) > len(all_headers):
            all_headers = headers
        all_rows.extend([n + 1] + row for row in rows)
        print(f"  {len(rows)} rows")
        time.sleep(1)  # be polite to the server
    return all_headers, all_rows


def save_excel(headers, rows, out_path):
    width = max(len(r) for r in rows) - 2  # minus Page and Links columns
    names = list(headers[:width]) + [""] * (width - len(headers))
    names = [h or f"Column {i + 1}" for i, h in enumerate(names)]
    columns = ["Page"] + names + ["Links"]

    wb = Workbook()
    ws = wb.active
    ws.title = "NAT data"
    ws.append(columns)
    for row in rows:
        page, values, links = row[0], row[1:-1], row[-1]
        values = values + [""] * (width - len(values))
        ws.append([page] + values + [links])

    # Drop the Links column if no row had any
    if not any(r[-1] for r in rows):
        ws.delete_cols(len(columns))
        columns.pop()

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, col in enumerate(ws.iter_cols(values_only=True), start=1):
        longest = max(len(str(v)) if v is not None else 0 for v in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(longest + 2, 8), 60)
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--pages", type=int, default=7, help="number of pages (default 7)")
    parser.add_argument("--out", default="dge_nat.xlsx", help="output Excel file")
    args = parser.parse_args()

    headers, rows = scrape(args.pages)
    if not rows:
        sys.exit("No data was extracted. The page layout may have changed.")
    save_excel(headers, rows, args.out)
    print(f"\nSaved {len(rows)} rows to {args.out}")


if __name__ == "__main__":
    main()
