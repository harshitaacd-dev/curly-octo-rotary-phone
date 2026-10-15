"""One sheet per year: States/UTs as rows, crime heads as columns.

Usage: python scripts/make_yearwise_tables.py data/clean data/clean/statewise_crimeheads_by_year.xlsx

Combines state_women_cyber_panel.csv and state_cyber_sexual_abusive.csv. Columns NCRB did not
publish for a year are left out of that year's sheet.
"""
import csv, os, sys
from collections import OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GROUPS = OrderedDict([
    ("Crimes against women", [
        ("women_total_cases", "Total crime against women"),
        ("female_population_lakh", "Female population (lakh)"),
        ("women_crime_rate", "Crime rate (per lakh women)"),
        ("women_chargesheeting_rate", "Women: chargesheeting rate (%)"),
        ("women_ipc_total", "Total IPC/BNS crimes"),
        ("women_sll_total", "Total SLL crimes"),
        ("rape", "Rape"),
        ("rape_women_18plus", "Rape: women 18+"),
        ("rape_girls_under18", "Rape: girls <18"),
        ("attempt_to_rape", "Attempt to rape"),
        ("murder_with_rape", "Murder with rape/gang rape"),
        ("sexual_intercourse_by_deceit_bns69", "Sexual intercourse by deceit (BNS 69)"),
        ("dowry_deaths", "Dowry deaths"),
        ("abetment_of_suicide_of_women", "Abetment of suicide of women"),
        ("acid_attack", "Acid attack"),
        ("attempt_acid_attack", "Attempt to acid attack"),
        ("acid_attack_incl_attempt", "Acid attack incl. attempt"),
        ("cruelty_by_husband_or_relatives", "Cruelty by husband or relatives"),
        ("kidnapping_abduction_of_women", "Kidnapping & abduction of women"),
        ("kidnapping_to_compel_marriage", "Kidnapping to compel marriage"),
        ("assault_outrage_modesty_354_incl_354a_d", "Assault on modesty (354 incl. 354A-D)"),
        ("assault_outrage_modesty_354_as_reported", "Assault on modesty (354, as reported)"),
        ("sexual_harassment_354a", "Sexual harassment (354A)"),
        ("assault_to_disrobe_354b", "Assault to disrobe (354B)"),
        ("voyeurism_354c", "Voyeurism (354C)"),
        ("stalking_354d", "Stalking (354D)"),
        ("insult_to_modesty_509", "Insult to modesty (509)"),
        ("human_trafficking_370", "Human trafficking (370)"),
        ("procuration_of_minor_girls", "Procuration of minor girls"),
        ("importation_of_girls", "Importation of girls"),
        ("selling_of_girls", "Selling of girls"),
        ("buying_of_girls", "Buying of girls"),
        ("dowry_prohibition_act", "Dowry Prohibition Act"),
        ("immoral_traffic_act", "Immoral Traffic (Prevention) Act"),
        ("domestic_violence_act", "Domestic Violence Act"),
        ("indecent_representation_act", "Indecent Representation of Women Act"),
        ("pocso_girl_victims", "POCSO (girl victims)"),
        ("women_centric_cyber_crimes", "Women-centric cyber crimes"),
        ("caw_sexually_explicit_67a_67b", "Women-centric: sexually explicit material (67A/67B)"),
    ]),
    ("Cyber crimes", [
        ("cyber_total_cases", "Total cyber crimes"),
        ("population_lakh", "Population (lakh)"),
        ("cyber_crime_rate", "Cyber crime rate (per lakh)"),
        ("cyber_chargesheeting_rate", "Cyber: chargesheeting rate (%)"),
        ("cyber_it_act_total", "Total under IT Act"),
        ("cyber_ipc_total", "Total under IPC/BNS"),
        ("cyber_sll_total", "Total under SLL"),
        ("cyber_fraud", "Fraud (2017+)"),
        ("cyber_criminal_breach_of_trust_fraud", "Criminal breach of trust/fraud (pre-2017)"),
        ("cyber_cheating", "Cheating"),
        ("cyber_forgery", "Forgery"),
        ("cyber_data_theft", "Data theft"),
        ("cyber_computer_related_offences_66", "Computer related offences (Sec 66)"),
        ("cyber_identity_theft_66c", "Identity theft (66C)"),
        ("cyber_cheating_by_personation_66d", "Cheating by personation (66D)"),
        ("cyber_fake_news", "Fake news on social media"),
        ("cyber_crimes_against_women", "Cyber crimes against women"),
        ("cyber_crimes_against_children", "Cyber crimes against children"),
    ]),
    ("Sexual / abusive cyber crimes", [
        ("sexual_abusive_total", "Sexual + abusive total"),
        ("sexual_abusive_share_pct", "Share of all cyber crimes (%)"),
        ("sexual_content_total", "Sexual content total"),
        ("abusive_total", "Abusive total"),
        ("it67_all_subsections", "Obscene/sexually explicit material (Sec 67, all)"),
        ("it67_obscene_material", "Obscene material (67)"),
        ("it67a_sexually_explicit_material", "Sexually explicit material (67A)"),
        ("it67b_child_sexual_abuse_material", "Child sexual abuse material (67B)"),
        ("it67_other_subsections", "Other Sec 67"),
        ("it67c_retention_by_intermediaries", "Retention by intermediaries (67C, not sexual)"),
        ("it66e_violation_of_privacy", "Violation of privacy (66E)"),
        ("cyber_stalking_bullying", "Cyber stalking/bullying"),
        ("cyber_blackmail_threatening", "Cyber blackmail/threatening"),
        ("defamation_morphing", "Defamation/morphing"),
        ("fake_profile", "Fake profile"),
        ("abetment_of_suicide_online", "Abetment of suicide (online)"),
        ("vs_women_stalking_bullying", "Vs women: stalking/bullying"),
        ("vs_women_blackmail_threatening", "Vs women: blackmail/threatening"),
        ("vs_women_pornography_obscene_material", "Vs women: pornography/obscene material"),
        ("vs_women_defamation_morphing", "Vs women: defamation/morphing"),
        ("vs_women_fake_profile", "Vs women: fake profile"),
        ("vs_women_sexual_intercourse_by_deceit", "Vs women: sexual intercourse by deceit"),
        ("vs_women_other", "Vs women: other"),
        ("vs_children_sexual_abuse_material", "Vs children: sexual abuse material"),
        ("vs_children_stalking_bullying", "Vs children: stalking/bullying"),
        ("vs_children_blackmail_harassment", "Vs children: blackmail/harassment"),
        ("vs_children_fake_profile", "Vs children: fake profile"),
        ("vs_children_online_games_gaming_apps", "Vs children: online games/gaming apps"),
        ("vs_children_other", "Vs children: other"),
    ]),
])
FILLS = {"Crimes against women": "F3D9E6", "Cyber crimes": "D6E4F0", "Sexual / abusive cyber crimes": "E2D9F0"}
THIN = Side(style="thin", color="BBBBBB")


def load(src):
    data = {}
    for fn in ("state_women_cyber_panel.csv", "state_cyber_sexual_abusive.csv"):
        for r in csv.DictReader(open(os.path.join(src, fn), encoding="utf8")):
            row = data.setdefault((r["state"], int(r["year"])), {"row_type": r["row_type"]})
            for k, v in r.items():
                if k not in ("state", "year", "row_type") and v != "":
                    row.setdefault(k, float(v))
    return data


def main(src, out):
    data = load(src)
    years = sorted({y for _, y in data})
    wb = Workbook()
    notes = wb.active
    notes.title = "Notes"
    lines = [
        "State/UT-wise crime heads by year, NCRB Crime in India 2012-2024",
        "",
        "One sheet per year. Rows: States/UTs, then Total (States), Total (UTs), All India.",
        "Columns: crime heads grouped as crimes against women, cyber crimes, and sexual/abusive cyber crimes.",
        "Values are cases registered, except rates (per lakh population), chargesheeting rate (%), population (lakh) and share (%).",
        "A column appears in a year's sheet only if NCRB published it for that year. An empty cell inside a sheet means not reported for that State/UT.",
        "'Vs women' / 'Vs children' columns are victim-based subsets of the offence columns: do not add them to the offence columns.",
        "Breaks: Sec 354A-D added 2013; POCSO girl victims in crime against women from 2017; IPC to BNS from 1 July 2024; cyber 'fraud' changes meaning in 2017.",
        "Boundaries: Telangana from 2014; Ladakh separate from 2020; D&N Haveli and Daman & Diu merged from 2020.",
        "Full codebook: data/clean/README.md.",
    ]
    for i, l in enumerate(lines, 1):
        notes.cell(i, 1, l).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
    notes.column_dimensions["A"].width = 130

    for y in years:
        states = [s for (s, yy), r in data.items() if yy == y and r["row_type"] == "state_ut"]
        aggs = [s for s in ("Total (States)", "Total (UTs)", "All India") if (s, y) in data]
        rows = sorted(states) + aggs
        cols = [(g, k, lab) for g, items in GROUPS.items() for k, lab in items
                if any(k in data[(s, y)] for s in rows)]
        ws = wb.create_sheet(str(y))
        ws.cell(1, 1, "State/UT").font = Font(bold=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
        # Group header row, then crime-head row.
        c = 2
        for g in GROUPS:
            n = sum(1 for gg, _, _ in cols if gg == g)
            if not n:
                continue
            cell = ws.cell(1, c, g)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")
            cell.fill = PatternFill("solid", fgColor=FILLS[g])
            if n > 1:
                ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + n - 1)
            c += n
        for j, (g, k, lab) in enumerate(cols, 2):
            cell = ws.cell(2, j, lab)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center")
            cell.fill = PatternFill("solid", fgColor=FILLS[g])
            cell.border = Border(bottom=THIN)
            ws.column_dimensions[get_column_letter(j)].width = 13
        ws.row_dimensions[2].height = 75
        for i, s in enumerate(rows, 3):
            name = ws.cell(i, 1, s)
            agg = s in aggs
            if agg:
                name.font = Font(bold=True)
            for j, (g, k, lab) in enumerate(cols, 2):
                v = data[(s, y)].get(k)
                if v is None:
                    continue
                cell = ws.cell(i, j, int(v) if v == int(v) and not k.endswith(("rate", "lakh", "pct")) else round(v, 2))
                cell.number_format = "#,##0" if isinstance(cell.value, int) else "0.0#"
                if agg:
                    cell.font = Font(bold=True)
            if s == "Total (States)":
                for j in range(1, len(cols) + 2):
                    ws.cell(i, j).border = Border(top=THIN)
        ws.column_dimensions["A"].width = 30
        ws.freeze_panes = "B3"
    wb.save(out)
    print("wrote", out, "sheets:", years, file=sys.stderr)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
