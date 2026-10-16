"""Bundle the cleaned CSVs into one Excel workbook.

Usage: python scripts/make_workbook.py data/clean data/clean/ncrb_women_cyber_2012_2024.xlsx
"""
import csv, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEETS = [
    ("State panel", "state_women_cyber_panel.csv",
     "Crimes against women and cyber crimes by State/UT and year: totals, rates and main crime heads"),
    ("Cyber sexual-abusive", "state_cyber_sexual_abusive.csv",
     "Cyber crimes that are sexual or abusive in nature, by State/UT and year"),
    ("National cyber heads", "national_cyber_crimeheads_2014_2023.csv",
     "All-India cyber crime cases by crime head, 2014-2023 (Lok Sabha USQ 438, 2 Dec 2025)"),
]
NOTES = [
    "Source: NCRB, Crime in India 2012-2024 (State/UT tables), plus Lok Sabha Unstarred Question 438 (MHA, 2 Dec 2025) for all-India 2015 cyber crime heads.",
    "Values are cases registered unless the column name says rate (per lakh population), pct or lakh.",
    "Empty cell = not published for that year at that level. It is not zero.",
    "Each year has one row per State/UT plus 'Total (States)', 'Total (UTs)' and 'All India' (see row_type). State/UT rows add up to All India wherever States have values.",
    "2015 cyber crime heads by State come from the second sheet of NCRB's 2015-edition cyber workbook.",
    "In 'Cyber sexual-abusive', add up only within one block: offence-based columns, vs_women_*, vs_children_*, caw_* (the last three are victim-based subsets).",
    "Breaks in definitions: 2013 law change added Sec 354A-D; POCSO girl victims count as crime against women from 2017; IPC to BNS from 1 July 2024 (use assault_outrage_modesty_354_incl_354a_d for trends); cyber 'fraud' changes meaning in 2017.",
    "Boundaries: Telangana from 2014; Ladakh separate from 2020; D&N Haveli and Daman & Diu merged from 2020.",
    "Full codebook: data/clean/README.md in the repository.",
]
HEAD_FILL = PatternFill("solid", fgColor="DDE5F0")


def value(v):
    if v == "":
        return None
    try:
        f = float(v)
        return int(f) if f == int(f) and "." not in v else f
    except ValueError:
        return v


def coverage_rows(path):
    rows = list(csv.DictReader(open(path, encoding="utf8")))
    years = sorted({int(r["year"]) for r in rows})
    out = []
    for col in rows[0]:
        if col in ("state", "year", "row_type"):
            continue
        have = {int(r["year"]) for r in rows if r[col] != "" and r["row_type"] == "state_ut"}
        india = {int(r["year"]) for r in rows if r[col] != "" and r["state"] == "All India"}
        out.append([col] + ["S" if y in have else ("N" if y in india else "") for y in years])
    return years, out


def main(src, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    ws["A1"] = "NCRB crimes against women and cyber crimes, 2012-2024"
    ws["A1"].font = Font(bold=True, size=14)
    r = 3
    for title, fn, desc in SHEETS:
        ws.cell(r, 1, title).font = Font(bold=True)
        ws.cell(r, 2, desc)
        r += 1
    ws.cell(r, 1, "Coverage").font = Font(bold=True)
    ws.cell(r, 2, "Which years each column has data: S = by State/UT, N = All India only")
    r += 2
    for n in NOTES:
        ws.cell(r, 1, "•")
        ws.cell(r, 2, n).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 120

    for title, fn, _ in SHEETS:
        rows = list(csv.reader(open(os.path.join(src, fn), encoding="utf8")))
        sh = wb.create_sheet(title)
        for i, row in enumerate(rows):
            sh.append(row if i == 0 else [value(v) for v in row])
        for c in range(1, len(rows[0]) + 1):
            cell = sh.cell(1, c)
            cell.font = Font(bold=True)
            cell.fill = HEAD_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            width = max(len(str(x[c - 1])) for x in rows[1:] if len(x) >= c) if len(rows) > 1 else 10
            sh.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 10), 45 if c == 1 else 16)
        sh.row_dimensions[1].height = 60
        sh.freeze_panes = "D2" if rows[0][:3] == ["state", "year", "row_type"] else "B2"
        sh.auto_filter.ref = sh.dimensions

    cov = wb.create_sheet("Coverage")
    r = 1
    for title, fn, _ in SHEETS[:2]:
        years, rows = coverage_rows(os.path.join(src, fn))
        cov.cell(r, 1, title).font = Font(bold=True)
        r += 1
        for c, h in enumerate(["column"] + years, 1):
            cell = cov.cell(r, c, h)
            cell.font = Font(bold=True)
            cell.fill = HEAD_FILL
        r += 1
        for row in rows:
            for c, v in enumerate(row, 1):
                cov.cell(r, c, v).alignment = Alignment(horizontal="left" if c == 1 else "center")
            r += 1
        r += 1
    cov.column_dimensions["A"].width = 45
    for c in range(2, 16):
        cov.column_dimensions[get_column_letter(c)].width = 6
    wb.save(out)
    print("wrote", out, file=sys.stderr)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
