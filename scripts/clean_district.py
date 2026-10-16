"""Turn NCRB district-wise Excel tables into tidy long CSVs.

Usage: python scripts/clean_district.py data/ncrb data/clean

For each topic in TOPICS, reads every year's district-wise workbook listed in
data/ncrb/catalog.csv and writes:
  <out>/district_<topic>_long.csv  one row per year x state x district x crime head
  <out>/district_<topic>_heads.csv every crime-head column found, per year
Districts are checked against each state's "Total District(s)" row; mismatches
are written to <out>/district_checks.csv.
"""
import csv, os, re, sys
from collections import defaultdict
import openpyxl

TOPICS = {
    "women": re.compile(r"district.?wise.*(crimes? (committed )?against women)", re.I),
    "cyber": re.compile(r"district.?wise.*cyber", re.I),
    "children": re.compile(r"district.?wise.*against children", re.I),
    "missing": re.compile(r"district.?wise.*missing", re.I),
}
STATE_ROW = re.compile(r"^\s*(State|UT|State/UT)\s*:\s*(.+?)\s*$", re.I)
TOTAL_ROW = re.compile(r"^\s*total", re.I)
HEADER_START = re.compile(r"^\s*(SL|S\.?\s*No\.?)\s*$", re.I)
COLNUM_ROW = re.compile(r"^\s*\[?1\]?\s*$")
# Police units reported alongside districts (railway police, CID, etc.).
SPECIAL_UNIT = re.compile(r"\bG\.?R\.?P\b|railway|\bCID\b|crime branch|\bSTF\b|\bEOW\b|\bcyber\b|"
                          r"\bSTATE\b.*\bUNIT\b|special|vigilance|\bCB\b|\bSOG\b|\bATS\b|\bPS\b", re.I)
# Fix the few state spellings NCRB itself gets wrong.
STATE_FIX = {"Andrha Pradesh": "Andhra Pradesh", "A & N Islands": "A&N Islands",
             "D&N Haveli": "D&N Haveli", "Delhi UT": "Delhi"}


def text(v):
    return "" if v is None else re.sub(r"\s+", " ", str(v).replace("​", "")).strip()


def number(v):
    s = text(v).replace(",", "")
    if s in ("", "-", "NA", "N.A.", "N.A"):
        return None
    try:
        f = float(s)
    except ValueError:
        return None
    return int(f) if f == int(f) else f


def filled_rows(path):
    """Read the first sheet, copying each merged cell's value across its range."""
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if r - 1 < len(rows) and c - 1 < len(rows[r - 1]):
                    rows[r - 1][c - 1] = v
    return rows


def parse(path):
    rows = filled_rows(path)
    h0 = next(i for i, r in enumerate(rows) if HEADER_START.match(text(r[0])))
    # Header ends at the column-number row (2017+) or the first State row (2014-16).
    h1 = h0 + 1
    while h1 < len(rows) and not STATE_ROW.match(text(rows[h1][0])) and not COLNUM_ROW.match(text(rows[h1][0])):
        h1 += 1
    data_start = h1 + 1 if COLNUM_ROW.match(text(rows[h1][0])) else h1
    heads = {}
    for c in range(2, len(rows[h0])):
        parts = []
        for r in range(h0, h1):
            t = text(rows[r][c])
            if t and (not parts or parts[-1] != t):
                parts.append(t)
        if parts and parts != ["YEAR"]:
            heads[c] = " > ".join(parts)
    records, totals, state = [], {}, None
    for r in rows[data_start:]:
        a, b = text(r[0]), text(r[1]) if len(r) > 1 else ""
        m = STATE_ROW.match(a)
        if m:
            state = STATE_FIX.get(m.group(2), m.group(2))
            continue
        if state is None or not b:
            continue
        vals = {c: number(r[c]) for c in heads if c < len(r)}
        if TOTAL_ROW.match(b):
            totals[state] = vals
        elif re.fullmatch(r"\[?\d+\]?", a) or a == "":
            records.append((state, b, vals))
    return heads, records, totals


def main(src, out):
    os.makedirs(out, exist_ok=True)
    catalog = list(csv.DictReader(open(os.path.join(src, "catalog.csv"), encoding="utf8")))
    checks = []
    for topic, rx in TOPICS.items():
        files = sorted({(int(r["year"]), r["file"]) for r in catalog
                        if rx.search(r["title"]) and r["format"] in ("xlsx", "xls")
                        and r["status"] in ("ok", "exists")})
        if not files:
            continue
        with open(os.path.join(out, f"district_{topic}_long.csv"), "w", newline="", encoding="utf8") as fl, \
             open(os.path.join(out, f"district_{topic}_heads.csv"), "w", newline="", encoding="utf8") as fh:
            wl, wh = csv.writer(fl), csv.writer(fh)
            wl.writerow(["year", "state", "district", "special_unit", "col", "crime_head", "value"])
            wh.writerow(["year", "col", "crime_head", "source_file"])
            for year, f in files:
                heads, records, totals = parse(os.path.join(src, f))
                for c, h in heads.items():
                    wh.writerow([year, c + 1, h, f])
                sums = defaultdict(lambda: defaultdict(float))
                for state, district, vals in records:
                    special = int(bool(SPECIAL_UNIT.search(district)))
                    for c, v in vals.items():
                        if v is not None:
                            wl.writerow([year, state, district, special, c + 1, heads[c], v])
                            sums[state][c] += v
                bad = [(s, c) for s, tv in totals.items() for c, v in tv.items()
                       if v is not None and abs(sums[s][c] - v) > 0.5]
                checks.append([topic, year, len(records), len(totals), len(heads), len(bad),
                               "; ".join(f"{s} col{c + 1}" for s, c in bad[:5])])
                print(topic, year, f"{len(records)} districts, {len(heads)} heads, "
                      f"{len(bad)} total mismatches", file=sys.stderr)
    with open(os.path.join(out, "district_checks.csv"), "w", newline="", encoding="utf8") as f:
        w = csv.writer(f)
        w.writerow(["topic", "year", "districts", "states", "crime_heads", "total_mismatches", "examples"])
        w.writerows(checks)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
