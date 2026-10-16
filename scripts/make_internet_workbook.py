"""Combine the internet-access datasets into one Excel workbook.

Usage: python scripts/make_internet_workbook.py data/clean data/clean/india_internet_access_by_state_age_gender.xlsx

Inputs (all in data/clean): internet_use_surveys_long.csv, state_internet_trai.csv,
india_internet_trai.csv, india_internet_worldbank.csv.
"""
import csv, os, sys
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD = PatternFill("solid", fgColor="1F4E78")
SUB = PatternFill("solid", fgColor="D9E1F2")
AGG = PatternFill("solid", fgColor="F2F2F2")
WHITE = Font(bold=True, color="FFFFFF")
BOLD = Font(bold=True)
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
GENDERS = ["Male", "Female", "Total"]
SECTORS = ["Total", "Rural", "Urban"]
SECTOR_LABEL = {"Total": "Rural + Urban", "Rural": "Rural", "Urban": "Urban"}
STATE_ORDER = [
    "Andhra Pradesh", "Arunachal Pradesh", "Assam", "Bihar", "Chhattisgarh", "Goa", "Gujarat",
    "Haryana", "Himachal Pradesh", "Jharkhand", "Karnataka", "Kerala", "Madhya Pradesh",
    "Maharashtra", "Manipur", "Meghalaya", "Mizoram", "Nagaland", "Odisha", "Punjab", "Rajasthan",
    "Sikkim", "Tamil Nadu", "Telangana", "Tripura", "Uttar Pradesh", "Uttarakhand", "West Bengal",
    "A&N Islands", "Chandigarh", "D&N Haveli and Daman & Diu", "Delhi", "Jammu & Kashmir", "Ladakh",
    "Lakshadweep", "Puducherry", "All India",
]
UTS = set(STATE_ORDER[28:36])


def read(path):
    with open(path) as f:
        return list(csv.DictReader(f))


def fnum(v):
    return float(v) if v not in ("", None) else None


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = note
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 45
    return 4


def header(ws, row, groups, first="State/UT"):
    """groups: list of (group label, [sub labels]). Writes a two-row header from `row`."""
    ws.cell(row, 1, first).fill = HEAD
    ws.cell(row, 1).font = WHITE
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=1)
    col = 2
    for label, subs in groups:
        c = ws.cell(row, col, label)
        c.fill, c.font, c.alignment = HEAD, WHITE, Alignment(horizontal="center", wrap_text=True)
        if len(subs) > 1:
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + len(subs) - 1)
        for i, s in enumerate(subs):
            c = ws.cell(row + 1, col + i, s)
            c.fill, c.font, c.alignment = SUB, BOLD, Alignment(horizontal="center", wrap_text=True)
        col += len(subs)
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row + 2, 2)
    return row + 2


def rows_out(ws, row, labels, values, fmt="0.0"):
    for label, vals in zip(labels, values):
        ws.cell(row, 1, label)
        for j, v in enumerate(vals):
            c = ws.cell(row, 2 + j, v)
            c.number_format = fmt
        if label in ("All India", "Total"):
            for j in range(len(vals) + 1):
                ws.cell(row, 1 + j).fill = AGG
                ws.cell(row, 1 + j).font = BOLD
        row += 1
    return row


def widths(ws, first=30, rest=11):
    ws.column_dimensions["A"].width = first
    for i in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(i)].width = rest


def survey_index(rows):
    idx = {}
    for r in rows:
        idx[(r["survey"], r["indicator"], r["state"], r["age_group"], r["gender"], r["sector"])] = fnum(r["value_pct"])
    return idx


def state_age_gender_sheet(wb, idx, indicator, name, note):
    ws = wb.create_sheet(name)
    row = title(ws, f"{indicator.capitalize()}, by State/UT, age and gender — CAMS 2022-23 (%)", note)
    ages = ["15-24", "15-29", "15+"]
    for sector in SECTORS:
        ws.cell(row, 1, SECTOR_LABEL[sector]).font = Font(bold=True, size=12)
        row += 1
        groups = [(f"Age {a}" if a != "15+" else "Age 15 and above", GENDERS) for a in ages]
        groups.append(("Gender gap, pp (male − female)", ages))
        row = header(ws, row, groups)
        vals = []
        for s in STATE_ORDER:
            v = [idx.get(("CAMS (NSS 79th round)", indicator, s, a, g, sector)) for a in ages for g in GENDERS]
            gaps = [round(v[i * 3] - v[i * 3 + 1], 1) if v[i * 3] is not None and v[i * 3 + 1] is not None
                    else None for i in range(3)]
            vals.append(v + gaps)
        row = rows_out(ws, row, STATE_ORDER, vals) + 2
    ws.freeze_panes = "B7"
    widths(ws)


def main(src, out):
    surveys = read(os.path.join(src, "internet_use_surveys_long.csv"))
    trai_state = read(os.path.join(src, "state_internet_trai.csv"))
    trai_nat = read(os.path.join(src, "india_internet_trai.csv"))
    wbank = read(os.path.join(src, "india_internet_worldbank.csv"))
    idx = survey_index(surveys)
    used, able = "used internet in last 3 months", "able to use internet"
    wb = Workbook()

    # README -------------------------------------------------------------------------------
    ws = wb.active
    ws.title = "README"
    lines = [
        ("India internet access by State/UT, age and gender, 2004–2026", True),
        ("", False),
        ("No single source gives State × age × gender every year. This workbook combines the official sources that do exist:", False),
        ("", False),
        ("Sheet", True),
        ("State_Age_Gender_Used — internet use in last 3 months, State/UT × age (15-24, 15-29, 15+) × gender × rural/urban. CAMS 2022-23 (NSS 79th round, MoSPI Report 590, Table 17). The only source with all three breakdowns at once.", False),
        ("State_Age_Gender_Able — same layout, ability to use the internet (Table 16).", False),
        ("State_Gender_Trend — State/UT × gender across surveys: NSS 2017-18, NFHS-5 2019-21, CAMS 2022-23; plus TRAI subscriptions per 100 people, Dec 2021–Jun 2026.", False),
        ("NFHS5_State_Gender — women and men age 15-49 who have ever used the internet, by State/UT. NFHS-5 2019-21, Table 3.7.", False),
        ("Age_Gender_National — all-India internet use by age and gender: NSS 2017-18 (Table 18), NFHS-5 (Table 3.6), CAMS 2022-23, CMS: Telecom 2025.", False),
        ("NSS2017_State_Gender — persons age 5+ who used internet in last 30 days / able to use it, State × gender × rural/urban. NSS 75th round, Report 585, Statements 7.5.1–7.6.1 (major States only).", False),
        ("TRAI_State_Subs — internet subscriptions per State/UT (million and per 100 people), Dec 2021–Jun 2026. Connections, not people; no age/gender.", False),
        ("TRAI_National — internet subscriptions every December 2004–2025 and June 2026.", False),
        ("WorldBank_National — % of population using the internet, 2004–2025 (ITU estimate).", False),
        ("Survey_Data_Long — every survey value in long format (survey, period, indicator, State, age, gender, sector) for filtering/pivots.", False),
        ("", False),
        ("Read before comparing", True),
        ("• Definitions differ: NSS 2017-18 = used in last 30 days (age 5+); NFHS-5 = ever used (age 15-49); CAMS 2022-23 and CMS 2025 = used in last 3 months. Compare levels within a survey, not across surveys.", False),
        ("• Age and gender by State together exist only in CAMS 2022-23. For other years only State × gender (NSS 2017-18, NFHS-5) or national age × gender are published. Finer cross-tabs need the unit-level data from microdata.gov.in / dhsprogram.com (registration required).", False),
        ("• NSS 2017-18 'Male' includes transgender persons. Rural Delhi is not shown separately. Small States/UTs are not published.", False),
        ("• TRAI counts connections (SIMs/broadband lines), so per-100 can exceed 100; it has no age or gender. TRAI's Dec 2021 State split misallocates Chhattisgarh, Jharkhand and some north-eastern States.", False),
        ("• Survey estimates for small States/UTs (Lakshadweep, Ladakh, A&N, Chandigarh) rest on small samples; see the reports' relative standard errors.", False),
        ("", False),
        ("Sources: MoSPI NSS Reports 585 (2017-18) and 590 (CAMS 2022-23); CMS: Telecom 2025 press note (29 May 2025); NFHS-5 India Report (IIPS/MoHFW, via dhsprogram.com FR375); TRAI quarterly Performance Indicator Reports; World Bank WDI (IT.NET.USER.ZS). PDFs/extracts and URLs in data/external/.", False),
        ("Built by scripts/make_internet_workbook.py from the CSVs in data/clean/.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(bold=True, size=13 if i == 1 else 11) if b else Font()
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 140

    # CAMS State x age x gender ----------------------------------------------------------
    cams_note = ("Percentage of persons who used the internet at least once in the 3 months before the survey "
                 "(any device, any place). Source: MoSPI, Comprehensive Annual Modular Survey 2022-23 "
                 "(NSS 79th round), Report 590, Table 17. Blocks: rural + urban, rural, urban.")
    state_age_gender_sheet(wb, idx, used, "State_Age_Gender_Used", cams_note)
    state_age_gender_sheet(wb, idx, able, "State_Age_Gender_Able",
                           cams_note.replace("who used the internet at least once in the 3 months before the "
                                             "survey (any device, any place)", "able to use the internet on the "
                                             "date of survey").replace("Table 17", "Table 16"))

    # State x gender trend ---------------------------------------------------------------
    ws = wb.create_sheet("State_Gender_Trend")
    row = title(ws, "Internet use by State/UT and gender across surveys (%), with TRAI subscriptions per 100",
                "Different definitions and age ranges in each survey (see column headers) — compare within a column "
                "group. TRAI = connections per 100 people (all ages, no gender split). Blank = not published.")
    trai_dates = sorted({r["as_of"] for r in trai_state})
    groups = [
        ("NSS 2017-18: age 5+, used in last 30 days", GENDERS),
        ("NFHS-5 2019-21: age 15-49, ever used", ["Male", "Female"]),
        ("CAMS 2022-23: age 15+, used in last 3 months", GENDERS),
        ("CAMS 2022-23: age 15-24, used in last 3 months", GENDERS),
        ("TRAI internet subscriptions per 100 people", [d[:7] for d in trai_dates]),
    ]
    row = header(ws, row, groups)
    tp = {(r["state"], r["as_of"]): fnum(r["subs_per100"]) for r in trai_state}
    vals = []
    for s in STATE_ORDER:
        v = [idx.get(("NSS 75th round", "used internet in last 30 days", s, "5+", g, "Total")) for g in GENDERS]
        v += [idx.get(("NFHS-5", "ever used internet", s, "15-49", g, "Total")) for g in ("Male", "Female")]
        v += [idx.get(("CAMS (NSS 79th round)", used, s, "15+", g, "Total")) for g in GENDERS]
        v += [idx.get(("CAMS (NSS 79th round)", used, s, "15-24", g, "Total")) for g in GENDERS]
        v += [tp.get((s, d)) for d in trai_dates]
        vals.append(v)
    rows_out(ws, row, STATE_ORDER, vals)
    widths(ws)

    # NFHS-5 State x gender --------------------------------------------------------------
    ws = wb.create_sheet("NFHS5_State_Gender")
    row = title(ws, "Women and men age 15-49 who have ever used the internet, by State/UT — NFHS-5 2019-21 (%)",
                "Source: National Family Health Survey (NFHS-5), 2019-21, India Report, Table 3.7.")
    row = header(ws, row, [("Ever used internet (%)", ["Men", "Women", "Gap, pp (men − women)"])])
    vals = []
    for s in STATE_ORDER:
        m = idx.get(("NFHS-5", "ever used internet", s, "15-49", "Male", "Total"))
        w = idx.get(("NFHS-5", "ever used internet", s, "15-49", "Female", "Total"))
        vals.append([m, w, round(m - w, 1) if m is not None and w is not None else None])
    rows_out(ws, row, STATE_ORDER, vals)
    widths(ws, rest=14)

    # National age x gender --------------------------------------------------------------
    ws = wb.create_sheet("Age_Gender_National")
    row = title(ws, "All-India internet use by age group and gender (%)",
                "Each block is one survey with its own definition. Rural/urban splits are in Survey_Data_Long.")
    blocks = [
        ("NSS 75th round 2017-18 — used internet in last 30 days (Report 585, Table 18)", "NSS 75th round",
         "used internet in last 30 days", ["5-14", "15-29", "15-59", "60+", "15+", "5+"], GENDERS),
        ("NSS 75th round 2017-18 — able to use internet (Report 585, Table 18)", "NSS 75th round",
         able, ["5-14", "15-29", "15-59", "60+", "15+", "5+"], GENDERS),
        ("NFHS-5 2019-21 — ever used internet (Table 3.6)", "NFHS-5", "ever used internet",
         ["15-19", "20-24", "25-29", "30-34", "35-39", "40-44", "45-49", "15-49"], ["Male", "Female"]),
        ("CAMS 2022-23 — used internet in last 3 months (Report 590, Table 17)", "CAMS (NSS 79th round)",
         used, ["15-24", "15-29", "15+"], GENDERS),
    ]
    for label, survey, ind, ages, genders in blocks:
        ws.cell(row, 1, label).font = Font(bold=True, size=12)
        row = header(ws, row + 1, [(SECTOR_LABEL[s], genders) for s in SECTORS], first="Age group")
        vals = [[idx.get((survey, ind, "All India", a, g, s)) for s in SECTORS for g in genders] for a in ages]
        row = rows_out(ws, row, [a.replace("+", " and above") for a in ages], vals) + 2
    ws.cell(row, 1, "CMS: Telecom Jan-Mar 2025 — used internet in last 3 months (press note, 29 May 2025)").font = Font(bold=True, size=12)
    row = header(ws, row + 1, [(SECTOR_LABEL[s], GENDERS) for s in SECTORS], first="Age group")
    vals = [[idx.get(("CMS: Telecom", used, "All India", a, g, s)) for s in SECTORS for g in GENDERS]
            for a in ("15-24", "15-29", "15+")]
    rows_out(ws, row, ["15-24", "15-29", "15 and above"], vals)
    ws.freeze_panes = None
    widths(ws, first=26)

    # NSS 2017-18 State x gender ---------------------------------------------------------
    ws = wb.create_sheet("NSS2017_State_Gender")
    row = title(ws, "Persons age 5+ who used the internet in the last 30 days / are able to use it, by State — NSS 2017-18 (%)",
                "Source: MoSPI, NSS 75th round (July 2017–June 2018), Report 585, Statements 7.6.1 and 7.5.1. "
                "Major States only. 'Male' includes transgender. Rural Delhi not published separately.")
    states = [s for s in STATE_ORDER if any(k[2] == s and k[0] == "NSS 75th round" for k in idx)]
    for ind in ("used internet in last 30 days", able):
        ws.cell(row, 1, ind.capitalize()).font = Font(bold=True, size=12)
        row = header(ws, row + 1, [(SECTOR_LABEL[s], GENDERS) for s in SECTORS])
        vals = [[idx.get(("NSS 75th round", ind, st, "5+", g, s)) for s in SECTORS for g in GENDERS] for st in states]
        row = rows_out(ws, row, states, vals) + 2
    ws.freeze_panes = None
    widths(ws)

    # TRAI State -------------------------------------------------------------------------
    ws = wb.create_sheet("TRAI_State_Subs")
    row = title(ws, "Internet subscriptions by State/UT — TRAI (million, and per 100 people)",
                "Source: TRAI quarterly Performance Indicator Reports (Dec quarter each year; Jun 2026 latest). "
                "Subscriptions, not people; no age or gender breakdown. Dec 2021 values for Chhattisgarh, "
                "Jharkhand and some NE States are misallocated by TRAI.")
    labels = [d[:7] for d in trai_dates]
    row = header(ws, row, [("Total subscriptions (million)", labels), ("Per 100 people", labels),
                           ("Rural per 100", labels), ("Urban per 100", labels)])
    by = {(r["state"], r["as_of"]): r for r in trai_state}
    vals = []
    for s in STATE_ORDER:
        v = []
        for col in ("internet_subs_mn", "subs_per100", "rural_subs_per100", "urban_subs_per100"):
            v += [fnum(by[(s, d)][col]) if (s, d) in by else None for d in trai_dates]
        vals.append(v)
    rows_out(ws, row, STATE_ORDER, vals, fmt="0.00")
    widths(ws)

    # TRAI national + World Bank ---------------------------------------------------------
    for name, data, note in (
        ("TRAI_National", trai_nat, "TRAI internet subscriptions (million; per 100 people). Up to Dec 2012 the total excludes "
         "internet on mobile phones — do not compare across the 2012/2013 break. Broadband threshold: 256 kbps to 2012, "
         "512 kbps 2013-22, 2 Mbps from 2023."),
        ("WorldBank_National", wbank, "World Bank WDI (ITU data): individuals using the internet (% of population) and related "
         "indicators. internet_users_mn = share × population (derived)."),
    ):
        ws = wb.create_sheet(name)
        row = title(ws, name.replace("_", " "), note)
        cols = list(data[0].keys())
        for j, c in enumerate(cols, 1):
            cell = ws.cell(row, j, c)
            cell.fill, cell.font, cell.alignment = HEAD, WHITE, Alignment(wrap_text=True)
        for i, r in enumerate(data, row + 1):
            for j, c in enumerate(cols, 1):
                v = r[c]
                try:
                    v = float(v) if c not in ("as_of", "year") else (int(v) if c == "year" else v)
                except ValueError:
                    pass
                ws.cell(i, j, None if v == "" else v)
        ws.freeze_panes = ws.cell(row + 1, 2)
        widths(ws, first=12, rest=14)

    # Long data --------------------------------------------------------------------------
    ws = wb.create_sheet("Survey_Data_Long")
    cols = list(surveys[0].keys())
    for j, c in enumerate(cols, 1):
        cell = ws.cell(1, j, c)
        cell.fill, cell.font = HEAD, WHITE
    for i, r in enumerate(surveys, 2):
        for j, c in enumerate(cols, 1):
            ws.cell(i, j, fnum(r[c]) if c == "value_pct" else r[c])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(surveys) + 1}"
    ws.freeze_panes = "A2"
    for j, w in enumerate([22, 14, 30, 28, 10, 9, 9, 10, 36, 42], 1):
        ws.column_dimensions[get_column_letter(j)].width = w

    wb.save(out)
    print(f"wrote {out}", file=sys.stderr)


if __name__ == "__main__":
    main(*sys.argv[1:3])
