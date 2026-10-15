"""Combine datasets/*.csv into one Excel workbook: datasets/india_it_rules_reports.xlsx"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
D = ROOT / "datasets"
OUT = D / "india_it_rules_reports.xlsx"
PLATFORMS = [("google", "Google"), ("meta", "Meta"), ("whatsapp", "WhatsApp"), ("x", "X"),
             ("snap", "Snapchat"), ("sharechat", "ShareChat"), ("moj", "Moj"),
             ("pinterest", "Pinterest"), ("roblox", "Roblox")]
F = "Arial"
HEAD = Font(name=F, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
BODY = Font(name=F, size=10)

frames = {}
for f, name in PLATFORMS:
    df = pd.read_csv(D / f"{f}.csv", dtype={"notes": str}).fillna({"notes": ""})
    df.insert(4, "month", df["period_end"].str[:7])
    frames[name] = df
alldata = pd.concat(frames.values(), ignore_index=True)

wb = Workbook()


def table_sheet(ws, df, widths=None):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for c in ws[1]:
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
    vcol = list(df.columns).index("value") + 1 if "value" in df.columns else None
    if vcol:
        for (c,) in ws.iter_rows(min_row=2, min_col=vcol, max_col=vcol):
            c.number_format = "#,##0.##"
    for i, col in enumerate(df.columns, 1):
        w = (widths or {}).get(col) or min(max(len(str(col)), int(df[col].astype(str).str.len().quantile(0.9)) if len(df) else 10) + 2, 60)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


# --- About -----------------------------------------------------------------
ws = wb.active
ws.title = "About"
about = [
    ("India IT Rules monthly compliance reports — data", True),
    ("Numbers extracted from the monthly reports that significant social media intermediaries publish under "
     "Rule 4(1)(d) of India's IT (Intermediary Guidelines and Digital Media Ethics Code) Rules, 2021.", False),
    ("Source reports and extraction scripts: github.com/harshitaacd-dev/curly-octo-rotary-phone (reports/, scripts/extract/).", False),
    ("", False),
    ("Sheets", True),
    ("Headline totals — key monthly figures per platform, calculated with SUMIFS from 'All data'.", False),
    ("All data — every extracted number from all platforms in one filterable table (one row per published number).", False),
    ("One sheet per platform — the same rows, split by platform.", False),
    ("Validation — each published total compared with the sum of its categories.", False),
    ("Report index — every archived report with the period it covers and its source.", False),
    ("", False),
    ("Columns", True),
    ("platform / service — who published it; service is the product (Facebook, Instagram, Pins…). '(global)' = worldwide figure, not India-only.", False),
    ("period_start / period_end / month — reporting period as stated in the report; 'month' = month the period ends.", False),
    ("section / category / metric — where in the report the number comes from, the reason or policy, and what is counted. category 'All' = the platform's own total.", False),
    ("value / unit — the number; unit is count, percent, minutes (median), or 'count (rounded in source)' for abbreviated values like 676.4K or 10.7L.", False),
    ("notes / source_file — corrections applied, and the archived report the row came from.", False),
    ("", False),
    ("Caveats", True),
    ("Platforms count different things (URLs, pieces of content, accounts, enforcements) — do not compare raw numbers across platforms.", False),
    ("Category names are kept as published and change over time as platforms revise their policies.", False),
    ("Where a platform's categories don't add up to its own total, the published numbers are kept; see the Validation sheet.", False),
    ("'Nil'/'None' are stored as 0; '-', 'N/A', 'Not Available' and '<1' cells are omitted.", False),
    ("Data collected 14–15 October 2026.", False),
]
for text, bold in about:
    ws.append([text])
    ws.cell(ws.max_row, 1).font = Font(name=F, bold=bold, size=14 if ws.max_row == 1 else (11 if bold else 10))
    ws.cell(ws.max_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 130

# --- Headline totals (formulas over All data) ------------------------------
SERIES = [  # (header, platform, service, section, metric, category or None = all categories)
    ("Google: complaints received", "Google", None, "user complaints", "complaints received", "All"),
    ("Google: removals on complaints", "Google", None, "user complaints", "removal actions", "All"),
    ("Google: removals by automated detection", "Google", None, "automated detection", "removal actions", "All"),
    ("Facebook: grievance reports received", "Meta", "Facebook", "user grievances", "reports received", "All"),
    ("Instagram: grievance reports received", "Meta", "Instagram", "user grievances", "reports received", "All"),
    ("Facebook: content actioned (approx.)", "Meta", "Facebook", "proactive enforcement", "content actioned", None),
    ("Instagram: content actioned (approx.)", "Meta", "Instagram", "proactive enforcement", "content actioned", None),
    ("WhatsApp: grievances received", "WhatsApp", None, "user grievances", "reports received", "All"),
    ("WhatsApp: accounts banned", "WhatsApp", None, "accounts banned", "accounts banned", "All"),
    ("X: grievances received (to Sep 2025)", "X", "X", "user grievances", "grievances received", None),
    ("X: reports received (from Oct 2025)", "X", "X", "user grievances", "reports received", "All"),
    ("Snapchat: content & account reports", "Snapchat", None, None, "content & account reports", "All"),
    ("ShareChat: user reports received", "ShareChat", None, "user reports", "user reports received", "All"),
    ("Moj: user reports received", "Moj", None, "user reports", "user reports received", "All"),
    ("Roblox: grievance reports received", "Roblox", "Roblox", "user grievances", "reports received", "All"),
]
months = sorted(alldata["month"].unique())
ws = wb.create_sheet("Headline totals")
n = len(alldata) + 1
rng = lambda col: f"'All data'!${col}$2:${col}${n}"
cols = {c: get_column_letter(i + 1) for i, c in enumerate(alldata.columns)}
ws.append(["month"] + [s[0] for s in SERIES])
for m in months:
    r = ws.max_row + 1
    ws.cell(r, 1, m)
    for j, (_, plat, svc, sec, met, cat) in enumerate(SERIES, 2):
        crit = [(cols["platform"], plat), (cols["metric"], met)]
        if svc: crit.append((cols["service"], svc))
        if sec: crit.append((cols["section"], sec))
        if cat: crit.append((cols["category"], cat))
        args = ",".join(f'{rng(c)},"{v}"' for c, v in crit)
        cnt = f'COUNTIFS({rng(cols["month"])},$A{r},{args})'
        ws.cell(r, j, f'=IF({cnt}=0,"",SUMIFS({rng(cols["value"])},{rng(cols["month"])},$A{r},{args}))')
        ws.cell(r, j).number_format = "#,##0"
for c in ws[1]:
    c.font, c.fill = HEAD, HEAD_FILL
    c.alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[1].height = 48
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.font = BODY
ws.column_dimensions["A"].width = 10
for j in range(2, len(SERIES) + 2):
    ws.column_dimensions[get_column_letter(j)].width = 16
ws.freeze_panes = "B2"
ws.cell(ws.max_row + 2, 1, "Blank = not reported that month. Month = month in which the reporting period ends "
        "(X's periods run 26th–25th). Content actioned (Meta) sums rounded per-policy figures, so it is approximate.").font = Font(name=F, italic=True, size=9)

# --- All data, per-platform, validation, index -------------------------------
table_sheet(wb.create_sheet("All data"), alldata)
for name, df in frames.items():
    table_sheet(wb.create_sheet(name), df)
table_sheet(wb.create_sheet("Validation"), pd.read_csv(D / "validation.csv"))
table_sheet(wb.create_sheet("Report index"), pd.read_csv(ROOT / "reports" / "index.csv").fillna(""))

wb.save(OUT)
print(OUT, len(alldata), "rows;", len(months), "months")
